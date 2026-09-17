#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка отчёта по практической работе № 1 (РОСА Фреш на VMware) в форматах DOCX и PDF.

Оформление — по «Краткой выписке из ГОСТ 7.32-2017»: Times New Roman 14 пт,
полуторный межстрочный интервал, выравнивание по ширине, абзацный отступ,
поля 30/15/20/20 мм, нумерация страниц внизу по центру (титульный лист
не нумеруется), подписи к рисункам по центру, название таблицы слева
над таблицей, каждый структурный элемент — с нового листа.

Скриншоты берутся из ../screenshots, готовые файлы складываются в ../report.
Требуется: python-docx, reportlab, Pillow (см. requirements.txt).
"""

import os
import sys

from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.dirname(BASE)
SHOTS = os.path.join(LAB, 'screenshots')
REPORT_DIR = os.path.join(LAB, 'report')
FONTS = os.path.join(BASE, 'fonts')

sys.path.insert(0, BASE)
import content  # noqa: E402

# ----------------------------------------------------------------- параметры страницы
# Вариант «ГОСТ 7.32-2017»: поля 30/15/20/20 мм, абзацный отступ 1,25 см.
# Если в методичке указаны поля 25/20/20/20 мм и отступ 1,5 см,
# замените значения ниже на 2.5, 2.0, 2.0, 2.0 и 1.5 соответственно.
PAGE = {
    'margin_left': 3.0,     # см
    'margin_right': 1.5,    # см
    'margin_top': 2.0,      # см
    'margin_bottom': 2.0,   # см
    'indent': 1.25,         # см, абзацный отступ
    'font_size': 14,        # пт, основной текст
    'leading': 21.0,        # пт, полуторный интервал для 14 пт
    'img_width': float(os.environ.get('IMG_WIDTH_CM', '15')),  # см, ширина рисунков
}

DOCX_NAME = 'Отчёт_ЛР1_Установка_РОСА_на_VMware.docx'
PDF_NAME = 'Отчёт_ЛР1_Установка_РОСА_на_VMware.pdf'


def img_size(path, width_cm):
    """Возвращает (ширина, высота) картинки в сантиметрах с сохранением пропорций."""
    with Image.open(path) as im:
        w, h = im.size
    return width_cm, width_cm * h / w


def in_toc(block):
    """Нужно ли включать структурный элемент (h1c) в содержание."""
    return len(block) < 3 or bool(block[2])


def skip_in_toc():
    """Тексты структурных элементов, которые не включаются в содержание."""
    return {b[1] for b in content.SECTIONS if b[0] == 'h1c' and not in_toc(b)}


def toc_entries(page_map):
    """Список (уровень, текст, номер страницы) для содержания."""
    out = []
    for block in content.SECTIONS:
        if block[0] == 'h1':
            out.append((0, block[1], page_map.get(block[1], '')))
        elif block[0] == 'h2':
            out.append((1, block[1], page_map.get(block[1], '')))
        elif block[0] == 'h1c' and in_toc(block):
            out.append((0, block[1], page_map.get(block[1], '')))
    return out


def counts_for(content_counts):
    """Числа для реферата: страницы, рисунки, таблицы, источники."""
    figs = sum(1 for b in content.SECTIONS if b[0] == 'fig')
    tables = sum(1 for b in content.SECTIONS if b[0] == 'table')
    return {
        'pages': content_counts.get('pages', ''),
        'figs': figs,
        'tables': tables,
        'sources': len(content.REFS),
    }


def format_text(text, counts):
    """Подставляет числа в текст (для реферата) и убирает экранирование."""
    try:
        return text.format(**counts)
    except (KeyError, IndexError):
        return text


# ----------------------------------------------------------------------------- DOCX

def build_docx(page_map, path, counts):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    TNR = 'Times New Roman'
    SIZE = PAGE['font_size']
    IND = PAGE['indent']
    BLANK = Pt(PAGE['leading'])   # «одна пустая строка»
    doc = Document()

    # --- страница, поля
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(PAGE['margin_left']), Cm(PAGE['margin_right'])
    sec.top_margin, sec.bottom_margin = Cm(PAGE['margin_top']), Cm(PAGE['margin_bottom'])
    sec.different_first_page_header_footer = True   # титульный лист без номера

    # --- базовый стиль
    normal = doc.styles['Normal']
    normal.font.name = TNR
    normal.font.size = Pt(SIZE)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), TNR)
    normal.element.rPr.rFonts.set(qn('w:cs'), TNR)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(IND)

    def set_spacing(style, pt):
        """Разрежение текста заголовка (межбуквенный интервал), pt в пунктах."""
        rpr = style.element.get_or_add_rPr()
        el = rpr.find(qn('w:spacing'))
        if el is None:
            el = OxmlElement('w:spacing')
            rpr.append(el)
        el.set(qn('w:val'), str(int(pt * 20)))

    # --- заголовки разделов (уровень 1) и подразделов (уровень 2, разреженный на 3 пт)
    for name, spacing in (('Heading 1', 0), ('Heading 2', 3)):
        st = doc.styles[name]
        st.font.name = TNR
        st.font.size = Pt(SIZE)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), TNR)
        st.element.rPr.rFonts.set(qn('w:cs'), TNR)
        p = st.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.first_line_indent = Cm(IND)
        p.space_before = Pt(0)
        p.space_after = BLANK
        p.line_spacing = 1.5
        p.keep_with_next = True
        if spacing:
            set_spacing(st, spacing)

    def make_style(name, **kw):
        st = doc.styles.add_style(name, 1)  # WD_STYLE_TYPE.PARAGRAPH
        st.base_style = doc.styles['Normal']
        f = st.font
        f.name = TNR
        f.size = Pt(kw.get('size', SIZE))
        f.bold = kw.get('bold', False)
        f.italic = kw.get('italic', False)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), TNR)
        st.element.rPr.rFonts.set(qn('w:cs'), TNR)
        p = st.paragraph_format
        p.alignment = kw.get('align', WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.first_line_indent = Cm(kw.get('indent', IND))
        p.left_indent = Cm(kw.get('left', 0))
        p.line_spacing = 1.5
        p.space_before = Pt(kw.get('before', 0))
        p.space_after = Pt(kw.get('after', 0))
        p.keep_with_next = kw.get('keep', False)
        return st

    cap_style = make_style('FigureCaption', align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                           before=6, after=PAGE['leading'], keep=True)
    tab_cap_style = make_style('TableCaption', align=WD_ALIGN_PARAGRAPH.LEFT, indent=0,
                               before=12, after=6, keep=True)
    list_style = make_style('ListGOST')
    center_style = make_style('CenterNoIndent', align=WD_ALIGN_PARAGRAPH.CENTER, indent=0)
    struct_style = make_style('StructHeading', align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                              bold=True, before=0, after=PAGE['leading'], keep=True)
    toc_style = make_style('TocGOST', align=WD_ALIGN_PARAGRAPH.LEFT, indent=0)

    # --- номер страницы внизу по центру (кроме титульного листа)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.line_spacing = 1.0
    run = footer.add_run()
    run.font.name = TNR
    run.font.size = Pt(SIZE)
    fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld); run._r.append(instr); run._r.append(fld2)

    # --- титульный лист
    for line in content.TITLE:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(line.get('gap', 0) or 6)
        r = p.add_run(line['t'])
        r.font.name = TNR
        r.font.size = Pt(line.get('sz', SIZE))
        r.bold = bool(line.get('b'))
        r.italic = bool(line.get('i'))

    fig_no = [0]

    def add_figure(fname, caption):
        fig_no[0] += 1
        path_img = os.path.join(SHOTS, fname)
        w_cm, h_cm = img_size(path_img, PAGE['img_width'])
        par = doc.add_paragraph(style='CenterNoIndent')
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path_img, width=Cm(w_cm), height=Cm(h_cm))
        cap = doc.add_paragraph(style='FigureCaption')
        cap.add_run('Рисунок %d – %s' % (fig_no[0], caption))

    def add_table(spec):
        cap = doc.add_paragraph(style='TableCaption')
        cap.add_run(spec['caption'])
        rows = [spec['head']] + spec['rows']
        table = doc.add_table(rows=len(rows), cols=len(spec['head']))
        table.style = 'Table Grid'
        table.autofit = False
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = table.cell(i, j)
                cp = cell.paragraphs[0]
                cp.paragraph_format.first_line_indent = Cm(0)
                cp.paragraph_format.line_spacing = 1.0
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(2)
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = cp.add_run(val)
                run.font.name = TNR
                run.font.size = Pt(12)
                run.bold = (i == 0)
        table.columns[0].width = Cm(6.0)
        table.columns[1].width = Cm(10.5)
        for row in table.rows:      # ширины ячеек фиксируются по столбцам
            row.cells[0].width = Cm(6.0)
            row.cells[1].width = Cm(10.5)
        doc.add_paragraph(style='CenterNoIndent').paragraph_format.space_after = Pt(0)

    for block in content.SECTIONS:
        kind = block[0]
        if kind == 'h1':
            doc.add_paragraph(block[1], style='Heading 1')
        elif kind == 'h2':
            doc.add_paragraph(block[1], style='Heading 2')
        elif kind == 'h1c':
            p = doc.add_paragraph(style='StructHeading')
            p.paragraph_format.page_break_before = True
            p.add_run(block[1])
        elif kind == 'p':
            doc.add_paragraph(format_text(block[1], counts), style='Normal')
        elif kind == 'list':
            for item in block[1]:
                doc.add_paragraph('– ' + item, style='ListGOST')
        elif kind == 'refs':
            for i, item in enumerate(block[1], 1):
                doc.add_paragraph('%d %s' % (i, item), style='ListGOST')
        elif kind == 'toc':
            for level, text, page in toc_entries(page_map):
                par = doc.add_paragraph(style='TocGOST')
                par.paragraph_format.left_indent = Cm(0.5 * level)
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.tab_stops.add_tab_stop(
                    Cm(21.0 - PAGE['margin_left'] - PAGE['margin_right'] - 0.5 * level), 2, 1)
                par.add_run('%s\t%s' % (text, page))
        elif kind == 'fig':
            add_figure(block[1], block[2])
        elif kind == 'table':
            add_table(block[1])

    doc.core_properties.title = ('Отчёт о выполнении практической работы № 1. '
                                 'Установка ОС РОСА на VMware Workstation')
    doc.core_properties.author = '____________'
    doc.save(path)
    return fig_no[0]


# ------------------------------------------------------------------------------ PDF

def build_pdf(path, counts):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as RLTTFont
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer,
                                    Image as RLImage, Table, TableStyle, PageBreak,
                                    NextPageTemplate, KeepTogether)
    from reportlab.platypus.tableofcontents import TableOfContents

    for name, fname in (('Tinos', 'Tinos-Regular.ttf'), ('Tinos-Bold', 'Tinos-Bold.ttf'),
                        ('Tinos-Italic', 'Tinos-Italic.ttf'), ('Tinos-BoldItalic', 'Tinos-BoldItalic.ttf')):
        pdfmetrics.registerFont(RLTTFont(name, os.path.join(FONTS, fname)))
    pdfmetrics.registerFontFamily('Tinos', normal='Tinos', bold='Tinos-Bold',
                                  italic='Tinos-Italic', boldItalic='Tinos-BoldItalic')

    SIZE = PAGE['font_size']
    LEAD = PAGE['leading']
    TOC_NUM_RESERVE = 1.2 * cm   # место под отточие и номер страницы в содержании
    IND = PAGE['indent'] * cm
    BLANK = LEAD

    body = ParagraphStyle('body', fontName='Tinos', fontSize=SIZE, leading=LEAD,
                          alignment=TA_JUSTIFY, firstLineIndent=IND)
    h1 = ParagraphStyle('h1', parent=body, fontName='Tinos-Bold', alignment=TA_LEFT,
                        spaceBefore=0, spaceAfter=BLANK, keepWithNext=1)
    h2 = ParagraphStyle('h2', parent=h1, charSpace=3, spaceBefore=0, spaceAfter=BLANK,
                        keepWithNext=1)
    h1c = ParagraphStyle('h1c', parent=body, fontName='Tinos-Bold', alignment=TA_CENTER,
                         firstLineIndent=0, spaceBefore=0, spaceAfter=BLANK, keepWithNext=1)
    cap = ParagraphStyle('cap', parent=body, alignment=TA_CENTER, firstLineIndent=0,
                         spaceBefore=6, spaceAfter=BLANK)
    tabcap = ParagraphStyle('tabcap', parent=body, alignment=TA_LEFT, firstLineIndent=0,
                            spaceBefore=12, spaceAfter=6, keepWithNext=1)
    li = ParagraphStyle('li', parent=body, firstLineIndent=IND)
    cover = ParagraphStyle('cover', parent=body, alignment=TA_CENTER, firstLineIndent=0,
                           leading=18)

    skip = skip_in_toc()

    class ReportDoc(BaseDocTemplate):
        def __init__(self, filename, **kw):
            BaseDocTemplate.__init__(self, filename, **kw)
            self.toc_pages = {}

        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph):
                name = flowable.style.name
                if name in ('h1', 'h2', 'h1c'):
                    text = flowable.getPlainText()
                    level = 1 if name == 'h2' else 0
                    self.toc_pages[text] = self.page
                    if text not in skip:
                        self.notify('TOCEntry', (level, text, self.page))

    left = PAGE['margin_left'] * cm
    right = PAGE['margin_right'] * cm
    top = PAGE['margin_top'] * cm
    bottom = PAGE['margin_bottom'] * cm
    frame = Frame(left, bottom, A4[0] - left - right, A4[1] - top - bottom, id='main',
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    def draw_page_number(canvas, doc_):
        canvas.saveState()
        canvas.setFont('Tinos', SIZE)
        canvas.drawCentredString(A4[0] / 2.0, 1.2 * cm, str(doc_.page))
        canvas.restoreState()

    doc = ReportDoc(path, pagesize=A4, leftMargin=left, rightMargin=right, topMargin=top,
                    bottomMargin=bottom,
                    title='Отчёт о выполнении практической работы № 1. Установка ОС РОСА на VMware',
                    author='____________')
    doc.addPageTemplates([
        PageTemplate(id='cover', frames=[frame], onPage=lambda c, d: None),
        PageTemplate(id='main', frames=[frame], onPage=draw_page_number),
    ])

    story = [NextPageTemplate('main')]
    for line in content.TITLE:
        story.append(Paragraph(line['t'] or '&nbsp;',
                               ParagraphStyle('c', parent=cover,
                                              fontName='Tinos-Bold' if line.get('b') else 'Tinos',
                                              fontSize=line.get('sz', SIZE),
                                              spaceAfter=(line.get('gap') or 6))))

    fig_no = [0]

    def add_figure(fname, caption):
        fig_no[0] += 1
        p = os.path.join(SHOTS, fname)
        w_cm, h_cm = img_size(p, PAGE['img_width'])
        story.append(KeepTogether([
            Spacer(1, 6),
            RLImage(p, width=w_cm * cm, height=h_cm * cm),
            Paragraph('Рисунок %d – %s' % (fig_no[0], caption), cap),
        ]))

    for block in content.SECTIONS:
        kind = block[0]
        if kind == 'h1':
            story.append(Paragraph(block[1], h1))
        elif kind == 'h2':
            story.append(Paragraph(block[1], h2))
        elif kind == 'h1c':
            story.append(PageBreak())
            story.append(Paragraph(block[1], h1c))
        elif kind == 'p':
            story.append(Paragraph(format_text(block[1], counts), body))
        elif kind == 'list':
            for item in block[1]:
                story.append(Paragraph('– ' + item, li))
        elif kind == 'refs':
            for i, item in enumerate(block[1], 1):
                story.append(Paragraph('%d %s' % (i, item), li))
        elif kind == 'toc':
            toc = TableOfContents()
            toc.dotsMinLevel = 0
            # rightIndent резервирует место под отточие и номер страницы: иначе у длинной
            # строки номера, занимающей всю полосу набора, номер уезжает за поле.
            toc.levelStyles = [
                ParagraphStyle('TOC1', fontName='Tinos', fontSize=SIZE, leading=LEAD,
                               firstLineIndent=0, rightIndent=TOC_NUM_RESERVE),
                ParagraphStyle('TOC2', fontName='Tinos', fontSize=SIZE, leading=LEAD,
                               leftIndent=0.5 * cm, firstLineIndent=0,
                               rightIndent=TOC_NUM_RESERVE),
            ]
            story.append(toc)
        elif kind == 'fig':
            add_figure(block[1], block[2])
        elif kind == 'table':
            spec = block[1]
            story.append(Paragraph(spec['caption'], tabcap))
            data = [spec['head']] + spec['rows']
            cell = ParagraphStyle('cell', fontName='Tinos', fontSize=12, leading=15,
                                  firstLineIndent=0)
            cellb = ParagraphStyle('cellb', parent=cell, fontName='Tinos-Bold')
            data = [[Paragraph(c, cellb if i == 0 else cell) for c in row]
                    for i, row in enumerate(data)]
            t = Table(data, colWidths=[6.0 * cm, 10.5 * cm], repeatRows=1)
            t.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, (0, 0, 0)),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            story.append(KeepTogether([t]))
            story.append(Spacer(1, BLANK))

    doc.multiBuild(story)
    return doc.toc_pages, doc.page


def main():
    os.makedirs(REPORT_DIR, exist_ok=True)
    pdf_path = os.path.join(REPORT_DIR, PDF_NAME)
    docx_path = os.path.join(REPORT_DIR, DOCX_NAME)

    counts = counts_for({'pages': 0})     # первая прикидка — для числа страниц
    pages_map, n_pages = {}, 0
    for _ in range(4):                    # число страниц уточняется итеративно
        pages_map, n_pages = build_pdf(pdf_path, counts)
        if counts['pages'] == n_pages:
            break
        counts = counts_for({'pages': n_pages})
    print('PDF: %s — %d с., %d рис., %d табл., %d ист.'
          % (pdf_path, n_pages, counts['figs'], counts['tables'], counts['sources']))

    n_figs_docx = build_docx(pages_map, docx_path, counts)
    print('DOCX: %s — %d рис.' % (docx_path, n_figs_docx))


if __name__ == '__main__':
    main()
